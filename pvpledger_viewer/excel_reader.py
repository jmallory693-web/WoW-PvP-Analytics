"""Read workbook evidence without evaluating formulas or modifying the XLSX."""
from datetime import date, datetime, time, timedelta
from decimal import Decimal
import hashlib
import io
import json
from pathlib import Path
import re

READER_VERSION = '1.0'
FIELDS = {'Date':'Date','Time':'Time','Battleground':'Map','Arena':'Map','Result':'Result',
          'Duration':'Duration (s)','Our score':'My team score','Enemy score':'Enemy team score',
          'KB':'KB','HK':'HK','Deaths':'Deaths','Damage (M)':'Damage','Healing (M)':'Healing',
          'Honor':'Honor (scoreboard)','Rating change':'Rating change','PvP item level':'PvP item level',
          'Notes':'Notes','Conquest':'Conquest awarded','Role':'Role','Team MMR':'Team MMR','Enemy MMR':'Enemy MMR'}
OBJECTIVES = {'Bases assaulted':'Bases Assaulted','Bases defended':'Bases Defended','Flag captures':'Flags Captured',
              'Flag returns':'Flags Returned','Carts controlled':'Carts Controlled',
              'Orb possessions':'Orb Possessions','Victory points':'Victory Points'}


def packed(value):
    if isinstance(value,datetime):return {'type':'datetime','value':value.isoformat()}
    if isinstance(value,date):return {'type':'date','value':value.isoformat()}
    if isinstance(value,time):return {'type':'time','value':value.isoformat()}
    if isinstance(value,timedelta):return {'type':'timedelta','seconds':value.total_seconds()}
    if isinstance(value,bytes):
        import base64
        return {'type':'bytes','base64':base64.b64encode(value).decode('ascii')}
    return value


def digest(value):return hashlib.sha256(json.dumps(value,sort_keys=True,ensure_ascii=False,default=packed).encode()).hexdigest()


def duration(value,number_format=''):
    if isinstance(value,timedelta):return value.total_seconds()
    if isinstance(value,time):return value.hour*3600+value.minute*60+value.second+value.microsecond/1e6
    if isinstance(value,(int,float)) and not isinstance(value,bool) and re.search(r'[hs]|\[m\]',number_format,re.I):return value*86400
    if isinstance(value,str) and re.fullmatch(r'\d+:\d{2}(?::\d{2})?',value.strip()):
        parts=list(map(int,value.split(':')))
        return parts[0]*60+parts[1] if len(parts)==2 else parts[0]*3600+parts[1]*60+parts[2]
    return None


def day(value):
    if isinstance(value,(datetime,date)):return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
    if isinstance(value,str):
        for pattern in ('%Y-%m-%d','%m/%d/%Y','%m/%d/%y'):
            try:return datetime.strptime(value.strip(),pattern).date().isoformat()
            except ValueError:pass
    return None


def clock(value):
    if isinstance(value,datetime):value=value.time()
    if isinstance(value,time):return value.isoformat()
    if isinstance(value,str):
        for pattern in ('%I:%M %p','%I:%M:%S %p','%H:%M','%H:%M:%S'):
            try:return datetime.strptime(value.strip(),pattern).time().isoformat()
            except ValueError:pass
    return None


def number(value):
    if isinstance(value,(int,float)) and not isinstance(value,bool):
        import math
        return value if math.isfinite(value) else None
    return None


def million_precision(value,fmt):
    digits=max(0,-Decimal(str(value)).as_tuple().exponent)
    match=re.search(r'\.([0#]+)',fmt)
    if match:digits=max(digits,len(match[1]))
    return float(Decimal(1000000)*Decimal(10)**(-digits)/2)


def read_workbook(path,raw=None,owners=None):
    import openpyxl
    path=Path(path);raw=raw if raw is not None else path.read_bytes()
    book=openpyxl.load_workbook(io.BytesIO(raw),data_only=False,keep_links=False)
    cached=openpyxl.load_workbook(io.BytesIO(raw),data_only=True,keep_links=False)
    # Prefixes belong to the user's workbook, never to application constants.
    # Roster tabs can inherit a prefix from a named summary/match-log title.
    owners = dict(owners or {})
    for sheet in book:
        title = str(sheet['A1'].value or '')
        match = re.match(r'^([^\s(]+)(?:\s*\([^)]*\))?\s+[\u2013\u2014\ufffd-]\s*(?:Battleground|Arena|Summary)',title,re.I)
        if match:
            prefix = sheet.title.split()[0]
            inferred = match.group(1)
            if prefix not in owners:
                owners[prefix] = inferred
    sha=hashlib.sha256(raw).hexdigest();sheets=[];matches=[];rosters=[];warnings=[]
    for sheet in book:
        cells=[];rows={}
        for values in sheet:
            row={}
            for cell in values:
                if cell.value is None and not cell.comment:continue
                value=cached[sheet.title][cell.coordinate].value if cell.data_type=='f' else cell.value
                record={'cell':cell.coordinate,'row':cell.row,'column':cell.column,'original':packed(cell.value),
                        'value':packed(value),'data_type':cell.data_type,'number_format':cell.number_format,
                        'formula':cell.value if cell.data_type=='f' else None,
                        'comment':cell.comment.text if cell.comment else None,
                        'hyperlink':cell.hyperlink.target if cell.hyperlink else None}
                cells.append(record);row[cell.column]=(cell,value,record)
                if cell.data_type=='e' or cell.data_type=='f' and value is None:
                    warnings.append({'sheet':sheet.title,'cell':cell.coordinate,'issue':'Excel error or unavailable formula cache; formula not executed'})
            if row:rows[values[0].row]=row
        header_row=next((n for n,r in rows.items() if {'Date','Result'}.issubset({str(v[1]) for v in r.values()}) and any(v[1] in ('Arena','Battleground') for v in r.values())),None)
        roster_header=next((n for n,r in rows.items() if {'Game #','Player','Realm'}.issubset({str(v[1]) for v in r.values()})),None)
        owner=owners.get(sheet.title.split()[0]);kind='match' if header_row else 'roster' if roster_header else 'context'
        sheets.append({'name':sheet.title,'state':sheet.sheet_state,'dimensions':[sheet.max_row,sheet.max_column],
                       'kind':kind,'owner':owner,'cells':cells,'merged':list(map(str,sheet.merged_cells.ranges)),
                       'tables':list(sheet.tables),'nonempty_rows':len(rows)})
        header_row=header_row or roster_header
        if not header_row:continue
        headers={col:str(v[1]).strip() for col,v in rows[header_row].items() if v[1] is not None}
        for n,row in rows.items():
            if n<=header_row:continue
            game=next((v[1] for col,v in row.items() if headers.get(col) in ('#','Game #')),None)
            if not isinstance(game,(int,float)) or isinstance(game,bool):continue
            fields={}
            for col,(cell,value,record) in row.items():
                label=headers.get(col,'Unlabelled '+cell.column_letter)
                if label in fields:label+=' ['+cell.coordinate+']'
                fields[label]=record
            if kind=='match' and not any(fields.get(k,{}).get('value') for k in ('Battleground','Arena','Date')):continue
            if kind=='roster' and not fields.get('Player',{}).get('value'):continue
            normalized={'Character':owner,'Realm':None};tolerances={};field_map={}
            for label,record in fields.items():
                cell,value,_=row[record['column']];target=FIELDS.get(label)
                if target=='Date':clean=day(value)
                elif target=='Time':clean=clock(value)
                elif target=='Duration (s)':clean=duration(value,cell.number_format)
                elif label.endswith(' (M)') and target in ('Damage','Healing'):
                    clean=int(Decimal(str(value))*1000000) if number(value) is not None else None
                    if clean is not None:tolerances[target]=million_precision(value,cell.number_format)
                elif target in ('KB','HK','Deaths','My team score','Enemy team score','Honor (scoreboard)','Conquest awarded','Rating change','PvP item level','Team MMR','Enemy MMR'):clean=number(value)
                else:clean=value if isinstance(value,(str,int,float,bool)) else None
                if target:normalized[target]=clean;field_map[label]=target
                if label in OBJECTIVES:
                    target='Objective: '+OBJECTIVES[label];normalized[target]=number(value);field_map[label]=target
            if kind=='match':
                bracket=str(fields.get('Bracket',{}).get('value') or '').casefold()
                normalized['Type']='Solo Shuffle' if 'shuffle' in bracket else 'Arena skirmish' if 'skirmish' in bracket else 'Rated arena' if 'rated' in bracket else None if 'Arena' in headers.values() else 'Battleground'
                normalized['Objectives']='; '.join(OBJECTIVES[k]+'='+str(v['value']) for k,v in fields.items() if k in OBJECTIVES and number(v['value']) is not None) or None
            else:
                normalized.update(Player=fields.get('Player',{}).get('value'),Realm=fields.get('Realm',{}).get('value') or '',Faction=fields.get('Faction',{}).get('value'))
            rid=digest([sha,sheet.title,n]);fingerprint=digest({k:v for k,v in normalized.items() if k not in ('Notes','Role')})
            entry={'id':rid,'sheet':sheet.title,'row':n,'game':game,'owner':owner,'kind':kind,'fields':fields,
                   'normalized':normalized,'field_map':field_map,'tolerances':tolerances,'fingerprint':fingerprint,'roster':[]}
            (matches if kind=='match' else rosters).append(entry)
    for match in matches:
        # Arena and battleground logs restart their own Game # sequences.
        # The supplied roster sheets refer to battleground logs only; the same
        # number on an arena row is not evidence of participant membership.
        match['roster']=[r for r in rosters if r['owner']==match['owner'] and r['game']==match['game']] if 'Battleground' in match['fields'] else []
    mark_duplicates(matches)
    document={'reader_version':READER_VERSION,'path':str(path.resolve()),'sha256':sha,'size':len(raw),'sheets':sheets,
              'matches':matches,'rosters':rosters,'warnings':warnings,'defined_names':list(book.defined_names),
              'row_count':sum(s['nonempty_rows'] for s in sheets),'formula_policy':'Original formula and cached value retained; no formulas or macros executed'}
    book.close();cached.close();return document


def mark_duplicates(matches):
    groups={}
    for match in matches:groups.setdefault(match['fingerprint'],[]).append(match['id'])
    for match in matches:match['duplicate_rows']=[r for r in groups[match['fingerprint']] if r!=match['id']]
